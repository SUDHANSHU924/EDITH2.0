from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel
from typing import List
from io import BytesIO
from pypdf import PdfReader
from groq import AsyncGroq
from core.config import settings
import os
import tempfile

router = APIRouter()

MAX_CHUNK_CHARS = 8000
CHUNK_OVERLAP = 400
MAX_CHUNKS = 12
MAX_FILE_CHARS = 200000


def _chunk_text(text: str) -> List[str]:
    chunks = []
    start = 0
    length = len(text)
    while start < length:
        end = min(start + MAX_CHUNK_CHARS, length)
        chunk = text[start:end]
        chunks.append(chunk)
        if end == length:
            break
        start = end - CHUNK_OVERLAP
    return chunks


async def _extract_text(upload: UploadFile) -> str:
    data = await upload.read()
    if not data:
        return ""

    content_type = upload.content_type or ""
    filename = (upload.filename or "").lower()

    if content_type == "application/pdf" or filename.endswith(".pdf"):
        try:
            reader = PdfReader(BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(pages)
        except Exception:
            return ""

    try:
        return data.decode("utf-8", errors="ignore")
    except Exception:
        return ""


async def _ask_groq(client: AsyncGroq, question: str, chunk: str, filename: str) -> str:
    response = await client.chat.completions.create(
        model=settings.GROQ_MODEL,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are EDITH File QA. Answer only from the provided file excerpt. "
                    "If the answer is not in the excerpt, say 'Not found in excerpt.'"
                ),
            },
            {
                "role": "user",
                "content": (
                    f"File: {filename}\nQuestion: {question}\n\nExcerpt:\n{chunk}"
                ),
            },
        ],
        max_tokens=512,
    )

    if response.choices and response.choices[0].message:
        return (response.choices[0].message.content or "").strip()
    return ""


async def _synthesize_answer(client: AsyncGroq, question: str, partials: List[str]) -> str:
    response = await client.chat.completions.create(
        model=settings.GROQ_MODEL,
        messages=[
            {
                "role": "system",
                "content": (
                    "Combine the partial answers into a single response. "
                    "Prefer concrete details and note when information is missing."
                ),
            },
            {
                "role": "user",
                "content": (
                    f"Question: {question}\n\nPartial answers:\n" + "\n\n".join(partials)
                ),
            },
        ],
        max_tokens=768,
    )

    if response.choices and response.choices[0].message:
        return (response.choices[0].message.content or "").strip()
    return ""

class FileRequest(BaseModel):
    type: str
    content: str
    filename: str = "document"


class PPTRequest(BaseModel):
    title: str
    slides: List[dict]
    theme: str = "dark"


class DocRequest(BaseModel):
    title: str
    content: str
    format: str = "docx"

@router.post("/create")
async def create(req: FileRequest):
    return {
        "status": "ok",
        "system": "04 - File Vault",
        "type": req.type,
        "message": f"File {req.filename}.{req.type} ready"
    }


@router.post("/create-ppt")
async def create_ppt(request: PPTRequest):
    """Generate PowerPoint from EDITH"""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)

    for slide_data in request.slides:
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)

        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(5, 5, 5)

        title_shape = slide.shapes.title
        if title_shape and slide_data.get("title"):
            title_shape.text = slide_data["title"]
            title_shape.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 240, 255)
            title_shape.text_frame.paragraphs[0].runs[0].font.size = Pt(36)

        if len(slide.placeholders) > 1:
            body = slide.placeholders[1]
            tf = body.text_frame
            for point in slide_data.get("points", []):
                p = tf.add_paragraph()
                p.text = point
                p.font.color.rgb = RGBColor(255, 255, 255)
                p.font.size = Pt(20)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
    prs.save(tmp.name)

    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        filename=f"{request.title}.pptx",
    )


@router.post("/generate")
async def generate_file(request: DocRequest):
    """AI-powered file generation"""
    groq_key = os.getenv("GROQ_API_KEY", "")
    if not groq_key:
        return {"error": "Add GROQ_API_KEY"}

    from groq import Groq

    client = Groq(api_key=groq_key)

    prompt = (
        f"Create a complete {request.format} document about: {request.title}\n\n"
        f"Content requirements: {request.content}\n\n"
        "Return structured content in JSON format:\n"
        "{\n"
        "    \"title\": \"...\",\n"
        "    \"sections\": [\n"
        "        {\"heading\": \"...\", \"content\": \"...\"}\n"
        "    ]\n"
        "}"
    )

    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[{"role": "user", "content": prompt}],
        max_tokens=2048,
    )

    return {
        "status": "ok",
        "content": response.choices[0].message.content,
        "format": request.format,
    }


@router.post("/qa")
async def qa(
    files: List[UploadFile] = File(...),
    question: str = Form("Summarize the files."),
):
    if not settings.GROQ_API_KEY:
        raise HTTPException(status_code=400, detail="GROQ_API_KEY not configured")

    client = AsyncGroq(api_key=settings.GROQ_API_KEY)
    question = (question or "Summarize the files.").strip()
    sources = []
    chunks: List[tuple[str, str]] = []

    for upload in files:
        filename = upload.filename or "file"
        text = (await _extract_text(upload)).strip()
        if not text:
            continue
        if len(text) > MAX_FILE_CHARS:
            text = text[:MAX_FILE_CHARS] + "\n[TRUNCATED]"
        sources.append(filename)
        for chunk in _chunk_text(text):
            chunks.append((filename, chunk))

    if not chunks:
        return {
            "answer": "No readable text found in the uploaded files.",
            "sources": list(dict.fromkeys(sources)),
        }

    if len(chunks) > MAX_CHUNKS:
        chunks = chunks[:MAX_CHUNKS]

    partials = []
    for filename, chunk in chunks:
        partial = await _ask_groq(client, question, chunk, filename)
        if partial:
            partials.append(f"[{filename}] {partial}")

    if not partials:
        return {
            "answer": "No relevant content found in the uploaded files.",
            "sources": list(dict.fromkeys(sources)),
        }

    if len(partials) == 1:
        final_answer = partials[0]
    else:
        final_answer = await _synthesize_answer(client, question, partials)

    return {
        "answer": final_answer or "No answer generated.",
        "sources": list(dict.fromkeys(sources)),
    }

@router.get("/status")
async def status():
    return {"status": "online", "system": "04"}
